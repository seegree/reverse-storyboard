from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image
import pytest

from edit_detector.document import write_document, format_timecode
from edit_detector.errors import AppError
from edit_detector.revision import parse_selection, revise_document


def test_select_51_shots_and_preserve_edits(tmp_path):
    images = []
    for n in range(51):
        path = tmp_path / f'{n}.jpg'
        Image.new('RGB', (32, 18), (n * 5, 40, 60)).save(path)
        images.append(path)
    source = write_document('movie.mp4', images, [f'Shot description {n}' for n in range(1, 52)], tmp_path / 'source.docx',
                            timecodes=[format_timecode(n) for n in range(51)])
    doc = Document(source)
    doc.paragraphs[6].runs[0].bold = True
    doc.save(source)
    original = source.read_bytes()
    selected = [3, 6, 11, 13, 14, 25, 34, 35, 38, 42]
    output = tmp_path / 'selected.docx'
    result = revise_document(source, '3 6 11 13 14 25 34 35 38 42', output)
    assert result.kept_numbers == selected
    assert result.original_count == 51
    assert source.read_bytes() == original
    doc = Document(output)
    captions = [p for p in doc.paragraphs if p._p.xpath('./w:pPr/w:numPr')]
    assert [p.text for p in captions] == [f'Shot description {n}\n{format_timecode(n - 1)}' for n in selected]
    assert captions[0].runs[0].bold
    ids = {p._p.pPr.numPr.numId.val for p in captions}
    assert len(ids) == 1
    numbering = doc.part.numbering_part.element
    num = numbering.num_having_numId(ids.pop())
    abstract = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{num.abstractNumId.val}"]')[0]
    assert abstract.find(qn('w:lvl')).find(qn('w:start')).get(qn('w:val')) == '1'
    with ZipFile(output) as z:
        media = [z.read(n) for n in z.namelist() if n.startswith('word/media/')]
        assert set(media) == {images[n - 1].read_bytes() for n in selected}
    second = tmp_path / 'second.docx'
    revise_document(output, '10, 1, 1', second)
    assert len(Document(second).inline_shapes) == 2


def test_legacy_shared_image_and_protections(tmp_path):
    image = tmp_path / 'image.jpg'
    Image.new('RGB', (20, 20), 'red').save(image)
    source = write_document('movie.mp4', [image] * 3, ['First shot', 'Second shot', 'Third shot'], tmp_path / 'source.docx')
    output = tmp_path / 'selected.docx'
    revise_document(source, '3 1', output)
    assert [p.text for p in Document(output).paragraphs][2:] == ['First shot', '', 'Third shot', '']
    assert len(Document(output).inline_shapes) == 2
    with ZipFile(output) as z:
        assert len([n for n in z.namelist() if n.startswith('word/media/')]) == 1
    with pytest.raises(AppError, match='new filename'):
        revise_document(source, '1', source, overwrite=True)
    with pytest.raises(AppError, match='already exists'):
        revise_document(source, '1', output)
    invalid = tmp_path / 'invalid.docx'
    Document().save(invalid)
    with pytest.raises(AppError, match='Choose an Edit Detector'):
        revise_document(invalid, '1', tmp_path / 'no.docx')


@pytest.mark.parametrize('selection', ['', '0', '-1', '52', '1-3', 'one'])
def test_invalid_selection(selection):
    with pytest.raises(AppError):
        parse_selection(selection, 51)


def test_timecode_rounding():
    assert format_timecode(59.9996) == '00:01:00.000'
    assert format_timecode(90061.123) == '25:01:01.123'
    assert parse_selection('3, 1\n3 2', 3) == [1, 2, 3]
