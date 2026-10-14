from pathlib import Path
import threading
from zipfile import ZipFile

from PIL import Image
import pytest

from edit_detector.caption import Caption
from edit_detector.errors import AppError, Cancelled
from edit_detector.pipeline import analyze
from edit_detector.video import Shot


@pytest.fixture
def staged_video(monkeypatch, tmp_path):
    video = tmp_path / "source.mp4"
    video.write_bytes(b"fixture")
    monkeypatch.setattr("edit_detector.pipeline.detect_shots", lambda *a, **kw: [Shot(1, 0, 24, 24), Shot(2, 24, 48, 24)])
    seen = []
    def extract(path, shots, folder, **kw):
        folder.mkdir()
        images = []
        for i, color in enumerate(("red", "blue"), 1):
            image = folder / f"{i}.jpg"
            Image.new("RGB", (320, 180), color).save(image)
            images.append(image)
        seen.extend(images)
        return images
    monkeypatch.setattr("edit_detector.pipeline.extract_frames", extract)
    return video, seen


def test_pipeline_order_embedding_and_cleanup(staged_video, tmp_path):
    video, frames = staged_video
    class FixtureCaptioner:
        def caption(self, image, **kw):
            return Caption(f"Fixture caption for image {image.stem}.", "", 0, 1)
    result = analyze(video, tmp_path / "result.docx", captioner=FixtureCaptioner())
    assert [r["number"] for r in result.shots] == [1, 2]
    assert [r["caption"]["text"] for r in result.shots] == ["Fixture caption for image 1.", "Fixture caption for image 2."]
    assert all(not p.exists() for p in frames)
    with ZipFile(result.output) as doc:
        assert len([n for n in doc.namelist() if n.startswith("word/media/")]) == 2


@pytest.mark.parametrize("failed_shots", [{2}, {1, 2}])
def test_caption_failures_keep_all_images_and_successful_captions(staged_video, tmp_path, failed_shots):
    from docx import Document
    video, frames = staged_video
    class FailingCaptioner:
        def caption(self, image, **kw):
            if int(image.stem) in failed_shots:
                raise AppError("model timed out")
            return Caption("Red cup", "Red cup", 0, 1)
    result = analyze(video, tmp_path / "result.docx", captioner=FailingCaptioner())
    doc = Document(result.output)
    assert len(doc.inline_shapes) == 2
    assert [s["number"] for s in result.caption_failures] == sorted(failed_shots)
    for n in (1, 2):
        text = doc.paragraphs[n * 2].text
        assert ("Caption unavailable" in text) == (n in failed_shots)
        if n not in failed_shots:
            assert text == "Red cup"
    assert all(not p.exists() for p in frames)


def test_cancel_exception_is_not_a_caption_failure(staged_video, tmp_path):
    video, frames = staged_video
    target = tmp_path / "existing.docx"
    target.write_bytes(b"previous document")
    class CancellingCaptioner:
        def caption(self, image, **kw):
            raise Cancelled("User cancelled")
    with pytest.raises(Cancelled):
        analyze(video, target, captioner=CancellingCaptioner(), overwrite=True)
    assert target.read_bytes() == b"previous document"
    assert all(not p.exists() for p in frames)


def test_cancel_during_caption_never_publishes_document(staged_video, tmp_path):
    video, frames = staged_video
    cancel = threading.Event()
    class CancellingCaptioner:
        def caption(self, image, **kw):
            cancel.set()
            return Caption("A fixture caption is returned.", "", 0, 1)
    with pytest.raises(Cancelled):
        analyze(video, tmp_path / "cancelled.docx", captioner=CancellingCaptioner(), cancel=cancel)
    assert not (tmp_path / "cancelled.docx").exists()
    assert all(not p.exists() for p in frames)


def test_optional_exports_and_rollback(staged_video, tmp_path, monkeypatch):
    import json
    from docx import Document
    video, frames = staged_video
    class FixtureCaptioner:
        def caption(self, image, **kw):
            return Caption('Red cup', '', 0, 1)
    folder = tmp_path / 'images'
    result = analyze(video, tmp_path / 'timed.docx', captioner=FixtureCaptioner(), include_timecodes=True, export_images=folder)
    assert result.images_folder == str(folder)
    assert len(list(folder.glob('*.jpg'))) == 2
    assert len(json.loads((folder / 'shots.json').read_text())['shots']) == 2
    assert '00:00:01.000 – 00:00:02.000' in Document(result.output).paragraphs[4].text
    with pytest.raises(AppError, match='already in use'):
        analyze(video, tmp_path / 'other.docx', captioner=FixtureCaptioner(), export_images=folder)
    def fail(*args, **kwargs):
        raise AppError('write failed')
    monkeypatch.setattr('edit_detector.pipeline.write_document', fail)
    rollback = tmp_path / 'rollback'
    with pytest.raises(AppError, match='write failed'):
        analyze(video, tmp_path / 'failed.docx', captioner=FixtureCaptioner(), export_images=rollback)
    assert not rollback.exists()
    assert not (tmp_path / 'failed.docx').exists()
