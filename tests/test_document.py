from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from docx import Document
from PIL import Image
import pytest

from edit_detector.document import write_document
from edit_detector.errors import AppError

NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}


def test_embedded_images_real_numbering_and_aspect_ratio(tmp_path):
    images = []
    for name, size, color in [("wide", (800, 200), "red"), ("tall", (200, 800), "blue")]:
        p = tmp_path / f"{name}.jpg"
        Image.new("RGB", size, color).save(p)
        images.append(p)
    captions = ["A red rectangle fills the frame.", "A blue rectangle fills the frame."]
    output = write_document("movie.mp4", images, captions, tmp_path / "result.docx")
    doc = Document(output)
    assert [p.text for p in doc.paragraphs] == ["Video Shot Breakdown", "movie.mp4", captions[0], "", captions[1], ""]
    assert len(doc.inline_shapes) == 2
    assert doc.inline_shapes[0].width / doc.inline_shapes[0].height == pytest.approx(4)
    assert doc.inline_shapes[1].width / doc.inline_shapes[1].height == pytest.approx(0.25)
    with ZipFile(output) as z:
        assert len([n for n in z.namelist() if n.startswith("word/media/")]) == 2
        tree = ET.fromstring(z.read("word/document.xml"))
        ids = tree.findall(".//w:numPr/w:numId", NS)
        assert len(ids) == 2 and ids[0].attrib == ids[1].attrib
        assert b'TargetMode="External"' not in z.read("word/_rels/document.xml.rels")
        num = ET.fromstring(z.read("word/numbering.xml"))
        assert num.findall(".//w:numFmt[@w:val='decimal']", NS)
    for p in images:
        p.unlink()
    assert len(Document(output).inline_shapes) == 2


def test_no_partial_document_or_overwrite(tmp_path):
    target = tmp_path / "existing.docx"
    target.write_bytes(b"original")
    with pytest.raises(AppError, match="already exists"):
        write_document("a.mp4", [tmp_path / "missing.jpg"], ["A caption."], target)
    assert target.read_bytes() == b"original"
    with pytest.raises((OSError, AppError)):
        write_document("a.mp4", [tmp_path / "missing.jpg"], ["A caption."], target, overwrite=True)
    assert target.read_bytes() == b"original"


def test_mismatched_caption_count(tmp_path):
    with pytest.raises(AppError, match="exactly one"):
        write_document("a.mp4", [Path("a.jpg")], [], tmp_path / "a.docx")
