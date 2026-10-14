"""Trim an existing app-generated DOCX without video, cached JPEGs or AI inference."""

from dataclasses import dataclass
from pathlib import Path
import re
from zipfile import BadZipFile

from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.exceptions import PackageNotFoundError
from docx.oxml.ns import qn

from .document import _numbering, save_document
from .errors import AppError
from .process import check_cancel


def parse_selection(text: str, count: int) -> list[int]:
    """Accept spaces/commas; keep unique positions in original document order."""
    tokens = [token for token in re.split(r"[\s,]+", text.strip()) if token]
    if not tokens or any(not re.fullmatch(r"[0-9]+", token) for token in tokens):
        raise AppError("Enter shot numbers separated by spaces or commas, for example: 3 6 11 13.")
    selected = sorted({int(token) for token in tokens})
    if selected[0] < 1 or selected[-1] > count:
        raise AppError(f"This document has {count} shots. Enter numbers from 1 to {count}.")
    return selected


@dataclass(frozen=True)
class RevisionResult:
    output: str
    kept_numbers: list[int]
    original_count: int


def revise_document(source: Path, selection: str, output: Path, *, overwrite=False, cancel=None) -> RevisionResult:
    source, output = source.expanduser().resolve(), output.expanduser().resolve()
    if source == output:
        raise AppError("Save the shortened document under a new filename; the source must be preserved.")
    check_cancel(cancel)
    try:
        document = Document(source)
    except (OSError, ValueError, KeyError, BadZipFile, PackageNotFoundError) as exc:
        raise AppError(f"Cannot open the shot document: {exc}") from exc
    body = document.element.body
    # Our old and new outputs both use a numbered caption followed by its image.
    groups = []
    current = None
    for element in list(body):
        if element.tag == qn("w:sectPr"):
            continue
        if element.xpath("./w:pPr/w:numPr/w:numId"):
            current = []
            groups.append(current)
        if current is not None:
            current.append(element)
    if (not groups or not document.paragraphs or document.paragraphs[0].text != "Video Shot Breakdown"
            or any(sum(len(e.xpath(".//pic:pic")) for e in group) != 1 for group in groups)):
        raise AppError("Choose an Edit Detector document with one numbered caption and image per shot.")
    selected = parse_selection(selection, len(groups))
    kept = set(selected)
    number_id = _numbering(document)
    for position, group in enumerate(groups, 1):
        check_cancel(cancel)
        if position not in kept:
            for element in group:
                body.remove(element)
        else:
            props = group[0].get_or_add_pPr().get_or_add_numPr()
            props.get_or_add_ilvl().val = 0
            props.get_or_add_numId().val = number_id
    # Remove unused image relationships so excluded pictures really leave the ZIP.
    # Shared images still referenced by a retained shot remain embedded.
    used = set(document.element.xpath(".//@r:embed")) | set(document.element.xpath(".//@r:link"))
    for rel in list(document.part.rels.values()):
        if rel.reltype == RT.IMAGE and rel.rId not in used:
            document.part.drop_rel(rel.rId)
    result = save_document(document, output, overwrite=overwrite, cancel=cancel)
    return RevisionResult(str(result), selected, len(groups))
