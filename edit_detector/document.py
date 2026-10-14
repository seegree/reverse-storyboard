"""Minimal editable Word output with genuine list numbering and embedded images."""

import os
from pathlib import Path
import tempfile
import math

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image

from .errors import AppError
from .process import check_cancel


def _numbering(document) -> int:
    root = document.part.numbering_part.element
    abstract_id = max((int(n.get(qn("w:abstractNumId"))) for n in root.findall(qn("w:abstractNum"))), default=-1) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl")
    lvl.set(qn("w:ilvl"), "0")
    for tag, value in (("start", "1"), ("numFmt", "decimal"), ("lvlText", "%1."), ("lvlJc", "left")):
        child = OxmlElement("w:" + tag)
        child.set(qn("w:val"), value)
        lvl.append(child)
    abstract.append(lvl)
    root.append(abstract)
    return root.add_num(abstract_id).numId


def write_document(video_name: str, images: list[Path], captions: list[str], output: Path,
                   *, overwrite=False, cancel=None, timecodes: list[str] | None = None) -> Path:
    if not images or len(images) != len(captions):
        raise AppError("Every shot must have exactly one image and one caption.")
    if timecodes is not None and len(timecodes) != len(images):
        raise AppError("Every shot must have exactly one timecode range.")
    output = output.expanduser().resolve()
    if output.suffix.lower() != ".docx":
        raise AppError("The output filename must end in .docx.")
    if output.exists() and not overwrite:
        raise AppError(f"Output already exists: {output}. Choose another name or use --overwrite.")
    check_cancel(cancel)
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    for name in ("Normal", "Title"):
        style = document.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in style.element.xpath("./w:pPr/w:pBdr"):
            border.getparent().remove(border)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Normal"].paragraph_format.line_spacing = 1.08
    document.styles["Title"].font.size = Pt(22)
    document.add_paragraph("Video Shot Breakdown", "Title")
    source = document.add_paragraph(video_name)
    source.paragraph_format.space_after = Pt(14)
    number_id = _numbering(document)
    for i, (image, caption) in enumerate(zip(images, captions), 1):
        check_cancel(cancel)
        paragraph = document.add_paragraph(caption)
        if timecodes is not None:
            run = paragraph.add_run("\n" + timecodes[i - 1])
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(90, 90, 90)
        props = paragraph._p.get_or_add_pPr().get_or_add_numPr()
        props.get_or_add_ilvl().val = 0
        props.get_or_add_numId().val = number_id
        fmt = paragraph.paragraph_format
        fmt.left_indent, fmt.first_line_indent = Inches(0.25), Inches(-0.25)
        fmt.keep_with_next, fmt.keep_together = True, True
        fmt.space_after = Pt(5)
        picture = document.add_paragraph()
        picture.paragraph_format.left_indent = Inches(0.25)
        picture.paragraph_format.space_after = Pt(16)
        picture.paragraph_format.keep_together = True
        with Image.open(image) as img:
            width, height = img.size
        scale = min(6.5 / width, 3.0 / height)
        shape = picture.add_run().add_picture(str(image), width=Inches(width * scale), height=Inches(height * scale))
        shape._inline.docPr.set("descr", f"Shot {i}: {caption}")
    document.core_properties.title = "Video Shot Breakdown"
    document.core_properties.subject = video_name
    document.core_properties.author = ""
    return save_document(document, output, overwrite=overwrite, cancel=cancel)


def format_timecode(seconds: float) -> str:
    """Elapsed video time, not embedded camera/SMPTE or drop-frame timecode."""
    if not math.isfinite(seconds) or seconds < 0:
        raise AppError("Timecodes require a nonnegative finite time.")
    milliseconds = int(seconds * 1000 + 0.5)
    hours, remainder = divmod(milliseconds, 3_600_000)
    minutes, remainder = divmod(remainder, 60_000)
    seconds, milliseconds = divmod(remainder, 1000)
    return f"{hours:02d}:{minutes:02d}:{seconds:02d}.{milliseconds:03d}"


def save_document(document, output: Path, *, overwrite=False, cancel=None) -> Path:
    output = output.expanduser().resolve()
    if output.suffix.lower() != ".docx":
        raise AppError("The output filename must end in .docx.")
    if output.exists() and not overwrite:
        raise AppError(f"Output already exists: {output}. Choose another name or use --overwrite.")
    output.parent.mkdir(parents=True, exist_ok=True)
    # Do not leave half a DOCX, or destroy a previous result if generation fails.
    fd, temporary = tempfile.mkstemp(prefix=".shot-document-", suffix=".docx", dir=output.parent)
    os.close(fd)
    try:
        document.save(temporary)
        check_cancel(cancel)
        if overwrite:
            os.replace(temporary, output)
        elif os.name == "nt":
            # Windows rename refuses an existing destination and works on FAT/exFAT.
            os.rename(temporary, output)
        else:
            # Hard-link publication is atomic and refuses a racing existing file.
            os.link(temporary, output)
        return output
    except FileExistsError as exc:
        raise AppError(f"Output already exists: {output}") from exc
    except OSError as exc:
        raise AppError(f"Cannot save Word document. Close it in Word and check the output folder: {exc}") from exc
    finally:
        Path(temporary).unlink(missing_ok=True)
