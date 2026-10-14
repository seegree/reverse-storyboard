"""Verify a portable Windows folder with child processes isolated from dev tools."""
import argparse
import json
import os
from pathlib import Path
import subprocess
import time
from zipfile import ZipFile

from docx import Document
from PIL import Image


def validate(bundle, folder, *, windows_only=True):
    if windows_only and os.name != 'nt':
        raise RuntimeError('Run portable validation on Windows.')
    bundle, folder = bundle.resolve(), folder.resolve()
    folder.mkdir(parents=True, exist_ok=False)
    cli = bundle / 'EditDetector-CLI.exe'
    env = os.environ.copy()
    env.pop('PYTHONPATH', None)
    env.pop('EDIT_DETECTOR_RUNTIME', None)
    env['PATH'] = os.pathsep.join([str(Path(env['SystemRoot']) / 'System32'), env['SystemRoot']])
    report = {'status': 'running', 'bundle': str(bundle), 'checks': []}

    def run(*args):
        start = time.monotonic()
        result = subprocess.run([str(cli), *map(str, args)], cwd=folder, env=env,
                                capture_output=True, text=True, errors='replace', timeout=300)
        report['checks'].append({'command': list(map(str, args)), 'exit_code': result.returncode,
                                 'seconds': round(time.monotonic() - start, 2),
                                 'stdout': result.stdout, 'stderr': result.stderr})
        if result.returncode:
            raise RuntimeError(f'Portable command failed: {args[0]}\n{result.stderr}')

    try:
        run('doctor')
        run('self-test', '--output', folder / 'self-test')
        video = folder / 'self-test' / 'three-shots.mp4'
        output = folder / 'Shots café.docx'
        images = folder / 'Images café'
        run('analyze', video, '-o', output, '--timecodes', '--images-dir', images,
            '--report', folder / 'analysis.json')
        doc = Document(output)
        assert len(doc.inline_shapes) == 3, 'Expected three embedded images'
        assert '00:00:' in doc.paragraphs[2].text, 'Missing timecodes'
        for image in images.glob('*.jpg'):
            with Image.open(image) as picture:
                picture.verify()
        assert len(list(images.glob('*.jpg'))) == 3, 'Expected three exported images'
        assert len(json.loads((images / 'shots.json').read_text())['shots']) == 3
        before = output.read_bytes()
        selected = folder / 'Selected shots.docx'
        run('revise', output, '--keep', '3 1', '-o', selected)
        assert output.read_bytes() == before, 'Revision changed the original document'
        assert len(Document(selected).inline_shapes) == 2
        with ZipFile(selected) as package:
            assert len([p for p in package.namelist() if p.startswith('word/media/')]) == 2
        report['status'] = 'passed'
    except Exception as exc:
        report['status'] = 'failed'
        report['error'] = str(exc)
        raise
    finally:
        (folder / 'validation.json').write_text(json.dumps(report, indent=2), encoding='utf-8')
    print(f'PASS: portable dependencies, inference, timecodes, exports, Unicode paths and revision. {folder}')


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('bundle', type=Path)
    parser.add_argument('--output', required=True, type=Path, help='New review folder outside the checkout')
    args = parser.parse_args()
    validate(args.bundle, args.output)
